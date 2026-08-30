from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOURCE = Path(r"F:\DESK\1FireDynamicsSimulation\SHENBBAO\FDS三类算例模型详细报告.md")
OUTPUT = Path(r"F:\DESK\1FireDynamicsSimulation\SHENBBAO\FDS三类算例模型详细报告_更新版.docx")


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_text(paragraph, text, size=10.5):
    # Minimal inline Markdown support for bold and code spans.
    token_re = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    for part in token_re.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=size - 0.5)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size)


def table_cells(line):
    return [value.strip() for value in line.strip().strip("|").split("|")]


def is_table_separator(line):
    return bool(re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line))


def create_table(document, rows):
    values = [table_cells(row) for row in rows]
    columns = max(len(row) for row in values)
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for index, row_values in enumerate(values):
        cells = table.add_row().cells
        for col in range(columns):
            cell = cells[col]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            text = row_values[col] if col < len(row_values) else ""
            add_text(paragraph, text, size=8.5)
            if index == 0:
                set_cell_shading(cell, "D9EAF7")
                for run in paragraph.runs:
                    run.bold = True
        if index == 0:
            set_repeat_table_header(table.rows[-1])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def set_normal_style(document):
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    for level, size, color in [(1, 16, (31, 78, 121)), (2, 14, (31, 78, 121)), (3, 12, (47, 84, 150))]:
        style = document.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)


def create_document(markdown):
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    set_normal_style(document)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("FDS 三类验证算例模型详细报告 | 2026-08-27")
    set_run_font(footer_run, size=8, color=(100, 100, 100))

    lines = markdown.splitlines()
    i = 0
    title_seen = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if re.match(r"^\|", stripped) and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [line]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i])
                i += 1
            create_table(document, rows)
            continue
        if re.match(r"^#{1,3}\s+", stripped):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and not title_seen:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(12)
                run = paragraph.add_run(text)
                set_run_font(run, size=20, bold=True, color=(31, 78, 121))
                title_seen = True
            else:
                paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
                add_text(paragraph, text, size={1: 16, 2: 14, 3: 12}[min(level, 3)])
            i += 1
            continue
        if stripped == "---":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            p_pr = paragraph._p.get_or_add_pPr()
            p_bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "B4C6E7")
            p_bdr.append(bottom)
            p_pr.append(p_bdr)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_text(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_text(paragraph, stripped[2:])
            i += 1
            continue
        if stripped.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(5)
            add_text(paragraph, stripped[2:], size=10)
            i += 1
            continue

        # Paragraphs are wrapped in the Markdown source; join adjacent text lines.
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            next_line = lines[i].strip()
            if not next_line or next_line == "---" or next_line.startswith("#") or next_line.startswith("|") or next_line.startswith("- ") or re.match(r"^\d+\.\s+", next_line):
                break
            paragraph_lines.append(next_line)
            i += 1
        paragraph = document.add_paragraph()
        add_text(paragraph, " ".join(paragraph_lines))

    return document


if __name__ == "__main__":
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    document = create_document(SOURCE.read_text(encoding="utf-8"))
    document.core_properties.title = "FDS 三类验证算例模型详细报告"
    document.core_properties.subject = "cabinet_01、FM_SNL_01/03 与 VTT_01 模型和边界条件"
    document.core_properties.author = "Codex"
    document.save(OUTPUT)
    print(OUTPUT)
